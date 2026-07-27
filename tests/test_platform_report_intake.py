from io import BytesIO

from openpyxl import Workbook
from fastapi.testclient import TestClient

import config
from core.platform_report_intake import parse_platform_report
from models.finance_ledger import FinanceLedger
from server.main import app


PROJECT_ID = "xinyu-hengtai-dakou"


def workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def keruyun_daily_brief() -> bytes:
    workbook = Workbook()
    info = workbook.active
    info.title = "制表信息"
    info.append(["筛选条件", None])
    info.append(["统计时间", "2026-07-17 ~ 2026-07-17"])
    info.append(["门店", "大口章鱼烧新余店"])

    operations = workbook.create_sheet("营业统计")
    operations.append(["营业数据", None])
    operations.append(["项目名称", "金额"])
    operations.append(["商品销售金额", 1410.30])
    operations.append(["附加费", 37.40])
    operations.append(["订单金额(无单收银+商品销售金额+附加费+附加消费税+收押金)", 1447.70])
    operations.append(["商户优惠", 133.12])
    operations.append(["订单配送支出", 59.40])
    operations.append(["服务费", 154.32])
    operations.append(["补贴", -16.12])
    operations.append(["营业收入(订单金额-商户优惠-订单配送支出+损益金额-服务费-收押金+补贴)", 1084.74])
    operations.append(["销售实收(营业收入-挂账)", 1084.74])
    operations.append(["销售收款(营业收入-储值支付-预付金抵扣-挂账)", 1084.74])

    payments = workbook.create_sheet("支付统计")
    payments.append(["支付方式收款统计（对应总收款）", None, None, None, None])
    payments.append(["支付方式", "支付笔数", "收款", "退款", "收款金额"])
    payments.append(["合计", 73, 71, 2, 1084.74])
    payments.append(["抖音团购券", 9, 9, 0, 115.24])
    payments.append(["美团团购券", 5, 5, 0, 80.02])
    payments.append(["现金", 8, 7, 1, 112.00])
    payments.append(["微信", 21, 21, 0, 371.00])
    payments.append(["支付宝", 8, 8, 0, 156.00])
    payments.append(["淘宝闪购餐饮", 13, 13, 0, 183.02])
    payments.append(["美团外卖", 9, 8, 1, 67.46])

    products = workbook.create_sheet("商品销售统计")
    products.append(["商品销售统计", None, None, None])
    products.append(["商品大类", "商品中类", "销售数量", "商品销售金额"])
    products.append(["合计", None, 75, 1410.30])
    products.append(["大口章鱼烧（加盟）", "经典必吃", 37, 584.00])
    return workbook_bytes(workbook)


def meituan_balance_statement() -> bytes:
    workbook = Workbook()
    balance = workbook.active
    balance.title = "余额流水"
    balance.append(["门店id", "日期", "类型", "金额(元)", "现有余额", "状态", "交易号"])
    balance.append([26267860, "2026-07-15 10:39:21", "2026.7.12~2026.7.14 账单", 222.28, 222.28, "交易成功", "41961952265"])
    balance.append([26267860, "2026-07-15 10:39:21", "余额提现", -222.28, 0, "交易成功", "3168625763"])
    workbook.create_sheet("推广费流水").append(["门店id", "日期", "类型", "金额(元)", "现有余额", "状态", "交易号"])
    return workbook_bytes(workbook)


def test_parse_keruyun_daily_brief_reads_all_relevant_sheets_and_signed_formula():
    report = parse_platform_report(
        "营业简报【大口章鱼烧新余店】2026-07-17~2026-07-17.xls",
        keruyun_daily_brief(),
    )

    assert report["report_type"] == "keruyun_daily_brief"
    assert report["period_start"] == "2026-07-17"
    assert report["period_end"] == "2026-07-17"
    assert report["summary"]["order_amount_minor"] == 144770
    assert report["summary"]["merchant_net_minor"] == 108474
    assert report["summary"]["subsidy_minor"] == -1612
    assert report["checks"]["source_formula_difference_minor"] == 0
    assert report["checks"]["payment_total_difference_minor"] == 0
    assert {item["method"] for item in report["payment_methods"]} == {
        "抖音团购券", "美团团购券", "现金", "微信", "支付宝", "淘宝闪购餐饮", "美团外卖",
    }


def test_parse_meituan_balance_statement_keeps_wallet_credit_and_withdrawal_out_of_revenue():
    report = parse_platform_report(
        "大口章鱼烧（新余恒太城店）2026-07-01~2026-07-18订单.xlsx",
        meituan_balance_statement(),
    )

    assert report["report_type"] == "meituan_balance_statement"
    assert report["summary"]["wallet_credit_minor"] == 22228
    assert report["summary"]["wallet_withdrawal_minor"] == 22228
    assert report["summary"]["revenue_impact_minor"] == 0
    assert report["rows"][0]["business_period_start"] == "2026-07-12"
    assert report["rows"][0]["business_period_end"] == "2026-07-14"
    assert report["rows"][1]["event_type"] == "wallet_withdrawal"


def test_unknown_multisheet_workbook_is_rejected_instead_of_guessed():
    workbook = Workbook()
    workbook.active.title = "未知表"
    workbook.active.append(["金额", 100])

    try:
        parse_platform_report("未知报表.xlsx", workbook_bytes(workbook))
    except ValueError as exc:
        assert "无法识别官方报表来源" in str(exc)
    else:
        raise AssertionError("unknown workbook must not be silently classified")


def test_keruyun_report_preview_then_confirm_is_idempotent(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PROJECT_DATA_DIR", tmp_path)
    client = TestClient(app)
    report_bytes = keruyun_daily_brief()

    preview = client.post(
        f"/api/projects/{PROJECT_ID}/finance/report-imports/preview",
        files={"file": ("营业简报.xls", report_bytes, "application/vnd.ms-excel")},
    )
    assert preview.status_code == 200
    assert preview.json()["report"]["summary"]["merchant_net_minor"] == 108474

    for _ in range(2):
        confirmed = client.post(
            f"/api/projects/{PROJECT_ID}/finance/report-imports/confirm",
            files={"file": ("营业简报.xls", report_bytes, "application/vnd.ms-excel")},
        )
        assert confirmed.status_code == 200
        assert confirmed.json()["report_type"] == "keruyun_daily_brief"

    ledger = FinanceLedger.for_project(PROJECT_ID)
    assert ledger.revenue_total(PROJECT_ID, "2026-07-17", "2026-07-17") == 108474
    assert ledger.merchant_net_total(PROJECT_ID, "2026-07-17", "2026-07-17") == 108474
    assert len(ledger.list_evidence_vouchers(PROJECT_ID, "2026-07-17", "2026-07-17")) == 1
    snapshot = client.get(
        f"/api/projects/{PROJECT_ID}/finance/daily-snapshot?date=2026-07-17",
    )
    assert snapshot.status_code == 200
    assert snapshot.json()["data_state"] == "available"
    assert snapshot.json()["day_activity"]["merchant_net_minor"] == 108474
    assert snapshot.json()["day_activity"]["accounting_revenue_minor"] == 108474
    close = client.post(
        f"/api/projects/{PROJECT_ID}/finance/daily-close/2026-07-17",
        json={
            "counted_cash": 112,
            "reserve_cash": 100,
            "merchant_net_confirmed": True,
            "fund_locations_reviewed": True,
            "outflows_reviewed": True,
        },
    )
    assert close.status_code == 200
    assert close.json()["status"] == "closed"
    assert close.json()["generated_reports"] == []
    assert client.get(
        f"/api/projects/{PROJECT_ID}/finance/overview?start=2026-07-17&end=2026-07-17",
    ).json()["profit_status"] == "incomplete"


def test_meituan_statement_confirm_creates_review_records_without_revenue(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PROJECT_DATA_DIR", tmp_path)
    ledger = FinanceLedger.for_project(PROJECT_ID)
    ledger.ensure_store(PROJECT_ID, "大口章鱼烧", "2026-07-01")
    ledger.upsert_fund_account(
        PROJECT_ID, account_key="meituan-wallet", name="美团商家钱包",
        account_kind="platform_wallet", owner_kind="platform",
        is_store_controlled=False, institution="美团", masked_number=None,
        opening_balance_minor=0, effective_from="2026-07-01",
        effective_to=None, status="active",
    )
    ledger.upsert_fund_account(
        PROJECT_ID, account_key="former-owner-bank", name="前老板代收银行卡",
        account_kind="bank", owner_kind="former_owner",
        is_store_controlled=False, institution="工商银行", masked_number=None,
        opening_balance_minor=0, effective_from="2026-07-01",
        effective_to=None, status="active",
    )
    revenue_before = ledger.revenue_total(PROJECT_ID, "2026-07-01", "2026-07-31")
    client = TestClient(app)
    report_bytes = meituan_balance_statement()

    response = client.post(
        f"/api/projects/{PROJECT_ID}/finance/report-imports/confirm",
        data={"wallet_account_key": "meituan-wallet", "destination_account_key": "former-owner-bank"},
        files={"file": ("美团余额流水.xlsx", report_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    assert response.json()["created_record_count"] == 2
    assert response.json()["revenue_impact_minor"] == 0
    assert ledger.revenue_total(PROJECT_ID, "2026-07-01", "2026-07-31") == revenue_before
    records = ledger.list_bookkeeping_records(PROJECT_ID, "2026-07-01", "2026-07-31")
    assert {item["transaction_kind"] for item in records} == {"platform_settlement", "account_transfer"}
    assert {item["status"] for item in records} == {"needs_review"}


def test_aggregated_platform_settlement_reconciles_against_each_business_day(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PROJECT_DATA_DIR", tmp_path)
    ledger = FinanceLedger.for_project(PROJECT_ID)
    ledger.ensure_store(PROJECT_ID, "大口章鱼烧", "2026-07-01")
    ledger.upsert_fund_account(
        PROJECT_ID, account_key="meituan-wallet", name="美团商家钱包",
        account_kind="platform_wallet", owner_kind="platform",
        is_store_controlled=False, institution="美团", masked_number=None,
        opening_balance_minor=0, effective_from="2026-07-01",
        effective_to=None, status="active",
    )
    daily_amounts = {
        "2026-07-12": 4791,
        "2026-07-13": 10169,
        "2026-07-14": 7268,
    }
    for business_date, amount_minor in daily_amounts.items():
        ledger.record_merchant_net_sale(
            PROJECT_ID,
            business_date=business_date,
            channel="美团外卖",
            amount_minor=amount_minor,
            source_basis="客如云营业简报·支付统计",
            evidence_status="confirmed",
            settlement_state="wallet_credited",
        )
    record_id = ledger.create_bookkeeping_record(
        PROJECT_ID,
        transaction_date="2026-07-15",
        direction="inflow",
        amount_minor=22228,
        transaction_kind="platform_settlement",
        business_scope="store",
        category_code=None,
        category_name="美团平台钱包账单",
        account_key="meituan-wallet",
        counterparty="美团",
        summary="2026.7.12~2026.7.14 账单",
        source_type="meituan_balance_statement",
        source_reference="batch-20260712-14",
        confidence="high",
        raw_data={
            "business_period_start": "2026-07-12",
            "business_period_end": "2026-07-14",
            "source_platform": "美团",
        },
        status="needs_review",
    )

    queue_item = next(
        item for item in ledger.reconciliation_queue(PROJECT_ID, "2026-07-12", "2026-07-15")
        if item["record"]["id"] == record_id
    )

    assert len(queue_item["candidates"]) == 3
    assert sum(item["suggested_match_minor"] for item in queue_item["candidates"]) == 22228
    for candidate in queue_item["candidates"]:
        ledger.match_reconciliation(
            PROJECT_ID,
            record_id,
            target_type=candidate["target_type"],
            target_id=candidate["id"],
            matched_amount_minor=candidate["suggested_match_minor"],
        )
    assert ledger.get_bookkeeping_record(PROJECT_ID, record_id)["reconciliation_status"] == "matched"

    try:
        ledger.match_reconciliation(
            PROJECT_ID,
            record_id,
            target_type=queue_item["candidates"][0]["target_type"],
            target_id=queue_item["candidates"][0]["id"],
            matched_amount_minor=1,
        )
    except Exception as exc:
        assert "已确认" in str(exc) or "剩余" in str(exc)
    else:
        raise AssertionError("fully reconciled sales cannot be matched again")
