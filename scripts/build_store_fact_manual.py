from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/store_manual/新余恒太城大口章鱼烧_门店经营事实手册_20260713.md"
OUTPUT = ROOT / "docs/store_manual/新余恒太城大口章鱼烧_门店经营事实手册_20260713.docx"
EVIDENCE = ROOT / "project_data/xinyu-hengtai-dakou/documents/evidence/2026-07-13"

NAVY = "17324D"
TEAL = "1B7F79"
PALE = "EAF4F2"
GRAY = "667085"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(section):
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Menlo"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(TEAL)
        else:
            paragraph.add_run(part)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline(p, value.strip())
            for run in p.runs:
                run.font.size = Pt(8.5)
                if r == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if r == 0:
                shade(cell, NAVY)
            elif r % 2 == 0:
                shade(cell, "F5F8FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    doc = Document()
    margins(doc.sections[0])
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "PingFang SC")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "PingFang SC")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string("243444")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color in (("Title", 28, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 10.5, NAVY)):
        style = styles[name]
        style.font.name = "PingFang SC"
        style._element.rPr.rFonts.set(qn("w:ascii"), "PingFang SC")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "PingFang SC")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("REAL STORE OPERATING RECORD")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph("新余恒太城五楼\n大口章鱼烧", style="Title")
    p = doc.add_paragraph()
    add_inline(p, "门店经营事实手册｜自 2026-07-01 接店以来的真实经营底账")
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_paragraph("A 已验证  ·  B 店主确认  ·  C 估算/待核")
    doc.add_paragraph("版本 2026-07-13 v1.1\n门店 ID  xinyu-hengtai-dakou\n事实截止  2026-07-12")
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    shade(box.cell(0, 0), PALE)
    box.cell(0, 0).text = "这不是展示型报告，而是收入、库存、成本、结算、人工与 Agent 自动化共同使用的事实底账。未知不填 0，计划不冒充发生，平台结算不重复计收入。"
    doc.add_section(WD_SECTION.NEW_PAGE)
    margins(doc.sections[-1])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    # Skip source title/version because cover already contains them.
    i = 4
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            table_rows = []
            while i < len(lines) and lines[i].startswith("|"):
                vals = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v) for v in vals):
                    table_rows.append(vals)
                i += 1
            if table_rows:
                add_table(doc, table_rows)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            shade_cell = doc.add_table(rows=1, cols=1)
            shade(shade_cell.cell(0, 0), PALE)
            add_inline(shade_cell.cell(0, 0).paragraphs[0], line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    doc.add_heading("附录｜原始图片缩略件", level=1)
    captions = {
        "8d4f74b98fefc1cbc1cdb92efd0740cd.jpg": "E01｜7 月 11 日客如云营业概况",
        "66c53fbc299cad0b92935c075e25b8c9.jpg": "E02｜7 月 12 日客如云营业概况",
        "59db595e85d6b78fe82163c4b928fecf.jpg": "E03｜门店现金盘点表",
        "6938d1c33a5b813c9cd44ae7f2dec358.jpg": "E04｜7 月混合银行流水",
        "b17520a087419a2029034163b9c75bbd.jpg": "E05｜2,570.80 元银行回单",
        "1ecf708a1083ac615d94ccf673ebc876.jpg": "E06｜7 月 5 日总部采购订单",
        "5b956d4ad892df3d39831940f08f5644.jpg": "E07｜7 月 6 日物料盘点",
        "f97fcba3c32f33d1db147729409964da.jpg": "E08｜线上补充采购（一）",
        "898cefe37073d4431161f4bce33b8f19.jpg": "E09｜线上补充采购（二）",
        "05c281a9b245612c0aaf63cd2fde2e81.jpg": "E10｜物料使用频率手写记录",
        "310db086fc2397b6c2dc1a8ddafe7428.jpg": "E11｜7 月 1—12 日门店收款统计",
        "2a3c471ff263798b67f2a9b59b86070d.jpg": "E12｜截至 7 月 12 日商品销售统计",
        "5789bd7f4b0b91c68c65562dbf154c58.jpg": "E13｜7 月 2 日本地供应商送货单",
        "ab914c6529daacf987642bea8ca835bf.jpg": "E14｜7 月 3 日本地供应商送货单",
        "42bf598dfbc0501b7ad6b07a21b62429.jpg": "E15｜7 月 5 日本地供应商送货单（含鸡蛋）",
        "3f29e4ee40d219f1debd8ddb1ac12df8.jpg": "E16｜7 月 7 日本地供应商送货单",
        "b187b516011ba77b8e86535c9ac83a08.jpg": "E17｜7 月 8 日本地供应商送货单",
        "28e881e21eae464f766b3f7068e42c14.jpg": "E18｜7 月 9 日本地供应商送货单",
        "25b12fd405016ee66e15494e10898c0b.jpg": "E19｜7 月 10 日本地供应商送货单",
        "8dcb7f403cfcd277b36df51470e01e54.jpg": "E20｜截至 7 月 12 日商品销售前列",
    }
    for name, caption in captions.items():
        doc.add_page_break()
        doc.add_heading(caption, level=2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        with Image.open(EVIDENCE / name) as im:
            ratio = im.width / im.height
        max_w, max_h = 15.5, 20.5
        width = min(max_w, max_h * ratio)
        p.add_run().add_picture(str(EVIDENCE / name), width=Cm(width))
        note = doc.add_paragraph("原图已按 SHA-256 固化；缩略件仅供定位，计算以正文确认口径和原图为准。")
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in note.runs:
            r.italic = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(GRAY)

    for section in doc.sections:
        margins(section)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("新余恒太城五楼大口章鱼烧｜经营事实手册｜2026-07-13")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
