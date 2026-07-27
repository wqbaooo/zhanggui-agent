"""Deterministic parsers for official store reports.

The platform adapters preserve the source's own signed fields and never turn a
wallet settlement or bank withdrawal into new sales revenue.  They only return
an inspectable preview; posting happens in the finance route after confirmation.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


def _text(value: Any) -> str:
    return str(value or "").strip()


def _minor(value: Any) -> int:
    if value in (None, ""):
        return 0
    text = _text(value).replace(",", "").replace("¥", "").replace("￥", "")
    text = re.sub(r"[^0-9.\-]", "", text)
    if text in {"", "-", ".", "-."}:
        return 0
    try:
        return int((Decimal(text) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except InvalidOperation:
        return 0


def _iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value).replace("年", "-").replace("月", "-").replace("日", "")
    match = re.search(r"(20\d{2})[./-](\d{1,2})[./-](\d{1,2})", text)
    if not match:
        return None
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def _period(value: Any) -> tuple[str | None, str | None]:
    dates = re.findall(r"(20\d{2})[./-](\d{1,2})[./-](\d{1,2})", _text(value))
    normalized = [f"{int(year):04d}-{int(month):02d}-{int(day):02d}" for year, month, day in dates]
    if not normalized:
        return None, None
    return normalized[0], normalized[-1]


def _rows(workbook: Any, sheet_name: str) -> list[list[Any]]:
    sheet = workbook[sheet_name]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _value_after_label(rows: list[list[Any]], label: str) -> Any:
    for row in rows:
        for index, value in enumerate(row):
            if _text(value).startswith(label):
                for candidate in row[index + 1 :]:
                    if candidate not in (None, ""):
                        return candidate
    return None


def _parse_keruyun(workbook: Any, filename: str) -> dict[str, Any]:
    info_rows = _rows(workbook, "制表信息")
    operation_rows = _rows(workbook, "营业统计")
    payment_rows = _rows(workbook, "支付统计")
    product_rows = _rows(workbook, "商品销售统计")

    period_start, period_end = _period(_value_after_label(info_rows, "统计时间"))
    store_name = _text(_value_after_label(info_rows, "门店"))

    summary = {
        "product_sales_minor": _minor(_value_after_label(operation_rows, "商品销售金额")),
        "surcharge_minor": _minor(_value_after_label(operation_rows, "附加费")),
        "order_amount_minor": _minor(_value_after_label(operation_rows, "订单金额")),
        "merchant_discount_minor": _minor(_value_after_label(operation_rows, "商户优惠")),
        "delivery_expense_minor": _minor(_value_after_label(operation_rows, "订单配送支出")),
        "service_fee_minor": _minor(_value_after_label(operation_rows, "服务费")),
        "subsidy_minor": _minor(_value_after_label(operation_rows, "补贴")),
        "merchant_net_minor": _minor(_value_after_label(operation_rows, "营业收入")),
        "sales_received_minor": _minor(_value_after_label(operation_rows, "销售收款")),
    }

    payment_methods: list[dict[str, Any]] = []
    header_index = next(
        (index for index, row in enumerate(payment_rows) if "支付方式" in {_text(value) for value in row}),
        None,
    )
    payment_total_minor = 0
    if header_index is not None:
        headers = [_text(value) for value in payment_rows[header_index]]
        for row in payment_rows[header_index + 1 :]:
            values = dict(zip(headers, row))
            method = _text(values.get("支付方式"))
            if not method:
                continue
            amount_minor = _minor(values.get("收款金额"))
            if method == "合计":
                payment_total_minor = amount_minor
                continue
            payment_methods.append(
                {
                    "method": method,
                    "payment_count": int(values.get("支付笔数") or 0),
                    "sales_count": int(values.get("收款") or 0),
                    "refund_count": int(values.get("退款") or 0),
                    "amount_minor": amount_minor,
                }
            )

    products: list[dict[str, Any]] = []
    product_header_index = next(
        (index for index, row in enumerate(product_rows) if "商品大类" in {_text(value) for value in row}),
        None,
    )
    if product_header_index is not None:
        headers = [_text(value) for value in product_rows[product_header_index]]
        for row in product_rows[product_header_index + 1 :]:
            values = dict(zip(headers, row))
            category = _text(values.get("商品大类"))
            if not category or category == "合计":
                continue
            products.append(
                {
                    "category": category,
                    "subcategory": _text(values.get("商品中类")),
                    "quantity": float(values.get("销售数量") or 0),
                    "sales_minor": _minor(values.get("商品销售金额")),
                }
            )

    computed_source_net = (
        summary["order_amount_minor"]
        - summary["merchant_discount_minor"]
        - summary["delivery_expense_minor"]
        - summary["service_fee_minor"]
        + summary["subsidy_minor"]
    )
    method_sum = sum(item["amount_minor"] for item in payment_methods)
    checks = {
        "source_formula_difference_minor": summary["merchant_net_minor"] - computed_source_net,
        "payment_total_difference_minor": summary["merchant_net_minor"] - method_sum,
        "reported_payment_total_difference_minor": summary["merchant_net_minor"] - payment_total_minor,
    }
    warnings = [label for label, value in checks.items() if value != 0]
    if period_start != period_end:
        warnings.append("not_single_business_day")
    if not payment_methods:
        warnings.append("missing_payment_breakdown")

    return {
        "report_type": "keruyun_daily_brief",
        "source_platform": "客如云",
        "filename": filename,
        "store_name": store_name,
        "period_start": period_start,
        "period_end": period_end,
        "summary": summary,
        "payment_methods": payment_methods,
        "products": products,
        "checks": checks,
        "warnings": warnings,
        "can_confirm": bool(period_start and period_start == period_end and summary["merchant_net_minor"] > 0 and not warnings),
        "accounting_effect": "daily_sales_and_receivable_position",
    }


def _parse_meituan(workbook: Any, filename: str) -> dict[str, Any]:
    statement_rows: list[dict[str, Any]] = []
    for sheet_name in ("余额流水", "推广费流水"):
        if sheet_name not in workbook.sheetnames:
            continue
        rows = _rows(workbook, sheet_name)
        if not rows:
            continue
        headers = [_text(value) for value in rows[0]]
        for raw_row in rows[1:]:
            values = dict(zip(headers, raw_row))
            occurred_on = _iso_date(values.get("日期"))
            transaction_type = _text(values.get("类型"))
            reference = _text(values.get("交易号"))
            if not occurred_on or not transaction_type or not reference:
                continue
            signed_amount_minor = _minor(values.get("金额(元)") if "金额(元)" in values else values.get("金额"))
            if sheet_name == "推广费流水":
                event_type = "promotion_fee"
            elif "提现" in transaction_type:
                event_type = "wallet_withdrawal"
            elif "账单" in transaction_type and signed_amount_minor >= 0:
                event_type = "wallet_credit"
            else:
                event_type = "unclassified_wallet_event"
            business_period_start, business_period_end = _period(transaction_type)
            statement_rows.append(
                {
                    "source_sheet": sheet_name,
                    "occurred_on": occurred_on,
                    "event_type": event_type,
                    "transaction_type": transaction_type,
                    "signed_amount_minor": signed_amount_minor,
                    "balance_minor": _minor(values.get("现有余额")),
                    "status": _text(values.get("状态")),
                    "reference": reference,
                    "business_period_start": business_period_start,
                    "business_period_end": business_period_end,
                }
            )

    dates = [row["occurred_on"] for row in statement_rows]
    credits = sum(row["signed_amount_minor"] for row in statement_rows if row["event_type"] == "wallet_credit")
    withdrawals = -sum(row["signed_amount_minor"] for row in statement_rows if row["event_type"] == "wallet_withdrawal")
    promotions = -sum(row["signed_amount_minor"] for row in statement_rows if row["event_type"] == "promotion_fee" and row["signed_amount_minor"] < 0)
    unknown_count = sum(row["event_type"] == "unclassified_wallet_event" for row in statement_rows)
    warnings = ["unclassified_wallet_events"] if unknown_count else []
    return {
        "report_type": "meituan_balance_statement",
        "source_platform": "美团",
        "filename": filename,
        "period_start": min(dates) if dates else None,
        "period_end": max(dates) if dates else None,
        "summary": {
            "wallet_credit_minor": credits,
            "wallet_withdrawal_minor": withdrawals,
            "promotion_fee_minor": promotions,
            "net_wallet_change_minor": sum(row["signed_amount_minor"] for row in statement_rows),
            "revenue_impact_minor": 0,
            "row_count": len(statement_rows),
        },
        "rows": statement_rows,
        "checks": {
            "duplicate_reference_count": len(statement_rows) - len({row["reference"] for row in statement_rows}),
            "unclassified_row_count": unknown_count,
        },
        "warnings": warnings,
        "can_confirm": bool(statement_rows and not warnings),
        "accounting_effect": "settlement_and_fund_movement_only",
    }


def parse_platform_report(filename: str, contents: bytes) -> dict[str, Any]:
    """Detect and parse an official report without guessing unknown schemas."""

    if not contents:
        raise ValueError("报表文件为空")
    try:
        # 客如云和美团导出的 OOXML 文件会把工作表 dimension 错写成 A1:A1。
        # openpyxl 的只读模式会因此只返回第一列；普通模式会按真实单元格读取。
        workbook = load_workbook(BytesIO(contents), read_only=False, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取报表，请确认文件是客如云或平台导出的 Excel") from exc
    try:
        sheet_names = set(workbook.sheetnames)
        if {"制表信息", "营业统计", "支付统计"}.issubset(sheet_names):
            return _parse_keruyun(workbook, filename)
        if "余额流水" in sheet_names:
            return _parse_meituan(workbook, filename)
        raise ValueError("无法识别官方报表来源；请勿用未知格式直接写账")
    finally:
        workbook.close()
