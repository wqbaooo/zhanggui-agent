"""Deterministic finance-file intake before any AI-assisted review.

The parser deliberately keeps bank money movement separate from revenue
recognition.  It extracts control totals and proposes classifications, but the
ledger only posts them after review.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from models.finance_ledger import money_to_minor


DATE_HEADERS = ("交易日期", "记账日期", "交易时间", "日期", "入账日期")
TIME_HEADERS = ("交易时间", "时间")
DIRECTION_HEADERS = ("收支类型", "收/支", "交易类型", "借贷标志", "收入/支出")
INFLOW_HEADERS = ("收入金额", "收入", "贷方发生额", "贷方金额", "转入金额")
OUTFLOW_HEADERS = ("支出金额", "支出", "借方发生额", "借方金额", "转出金额")
AMOUNT_HEADERS = ("交易金额", "金额", "发生额")
COUNTERPARTY_HEADERS = ("对方户名", "交易对方", "对方名称", "商户名称", "收款方", "付款方")
SUMMARY_HEADERS = ("摘要", "交易摘要", "商品说明", "备注", "交易说明", "用途")
REFERENCE_HEADERS = ("交易流水号", "流水号", "交易单号", "订单号", "参考号", "商户单号")
BALANCE_HEADERS = ("余额", "账户余额", "交易后余额")


def _first(row: dict[str, Any], headers: Iterable[str]) -> str:
    for header in headers:
        value = row.get(header)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _amount(value: Any) -> int:
    text = str(value or "").strip().replace(",", "").replace("¥", "").replace("￥", "")
    if not text or text in {"-", "--"}:
        return 0
    text = re.sub(r"[^0-9.()\-]", "", text)
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    try:
        return abs(money_to_minor(text))
    except Exception:
        return 0


def _normalize_date(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    text = text.replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-").replace(".", "-")
    match = re.search(r"(20\d{2})-(\d{1,2})-(\d{1,2})", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    for fmt in ("%Y%m%d", "%m-%d-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def _direction(row: dict[str, Any]) -> tuple[str, int]:
    inflow = _amount(_first(row, INFLOW_HEADERS))
    outflow = _amount(_first(row, OUTFLOW_HEADERS))
    label = _first(row, DIRECTION_HEADERS).lower()
    generic = _first(row, AMOUNT_HEADERS)
    generic_minor = _amount(generic)
    if inflow:
        return "inflow", inflow
    if outflow:
        return "outflow", outflow
    if any(token in label for token in ("收入", "转入", "贷", "收款", "入账")):
        return "inflow", generic_minor
    if any(token in label for token in ("支出", "转出", "借", "付款", "消费")):
        return "outflow", generic_minor
    if str(generic).strip().startswith("-"):
        return "outflow", generic_minor
    return "unknown", generic_minor


def classify_finance_transaction(
    *,
    direction: str,
    counterparty: str,
    summary: str,
    account_owner_kind: str,
) -> dict[str, Any]:
    text = f"{counterparty} {summary}".lower()
    result = {
        "transaction_kind": "unclassified",
        "business_scope": "store" if account_owner_kind == "store" else "unknown",
        "category_code": None,
        "category_name": "待归类",
        "confidence": "low",
        "classification_reason": "未找到足够可靠的分类依据，需人工确认",
    }

    if any(token in text for token in ("个人消费", "家庭", "私人")):
        return {**result, "transaction_kind": "personal_spending", "business_scope": "personal", "category_name": "个人消费", "confidence": "high", "classification_reason": "摘要明确标记为个人或家庭支出"}
    if direction == "inflow" and any(token in text for token in ("美团", "淘宝闪购", "饿了么", "京东", "抖音", "平台结算", "商户结算")):
        return {**result, "transaction_kind": "platform_settlement", "business_scope": "store", "category_name": "平台结算到账", "confidence": "high", "classification_reason": "交易对方或摘要命中平台结算；仅证明资金到账，不自动确认收入"}
    if direction == "inflow" and any(token in text for token in ("前老板", "代收款", "转回代收")):
        return {**result, "transaction_kind": "former_owner_transfer", "business_scope": "store", "category_name": "前老板转回代收款", "confidence": "high", "classification_reason": "摘要明确为前老板代收款转回，不重复增加收入"}
    if any(token in text for token in ("借款", "贷款放款")):
        kind = "loan_in" if direction == "inflow" else "loan_repayment"
        return {**result, "transaction_kind": kind, "business_scope": "store", "category_code": "2003", "category_name": "借款" if direction == "inflow" else "偿还借款本金", "confidence": "high", "classification_reason": "摘要明确包含借款或贷款"}
    if direction == "outflow" and any(token in text for token in ("工资", "薪资", "薪酬")):
        return {**result, "transaction_kind": "operating_expense", "business_scope": "store", "category_code": "6001", "category_name": "员工工资", "confidence": "high", "classification_reason": "摘要命中工资或薪酬"}
    if direction == "outflow" and any(token in text for token in ("房租", "租金", "商场管理费")):
        return {**result, "transaction_kind": "operating_expense", "business_scope": "store", "category_code": "6002", "category_name": "房租与商场费", "confidence": "high", "classification_reason": "摘要命中房租或商场费用"}
    if direction == "outflow" and any(token in text for token in ("电费", "水费", "燃气", "物业")):
        return {**result, "transaction_kind": "operating_expense", "business_scope": "store", "category_code": "6003", "category_name": "水电燃气", "confidence": "high", "classification_reason": "摘要命中水电、燃气或物业费用"}
    if direction == "outflow" and any(token in text for token in ("包装", "餐盒", "打包盒", "纸袋")):
        return {**result, "transaction_kind": "inventory_purchase", "business_scope": "store", "category_code": "1402", "category_name": "包装耗材库存", "confidence": "high", "classification_reason": "采购包装耗材先进入库存，不直接计入当期成本"}
    if direction == "outflow" and any(token in text for token in ("采购", "菜市场", "蔬菜", "食材", "原料", "供应商")):
        return {**result, "transaction_kind": "inventory_purchase", "business_scope": "store", "category_code": "1401", "category_name": "原材料库存", "confidence": "high", "classification_reason": "采购食材或原料先进入库存，不直接计入当期成本"}
    if any(token in text for token in ("提现", "转账", "内部转账")):
        return {**result, "transaction_kind": "account_transfer", "category_name": "账户间转账", "confidence": "medium", "classification_reason": "摘要显示为资金位置转移，需确认转入或转出账户"}
    return result


def _rows_from_csv(contents: bytes) -> list[dict[str, Any]]:
    decoded = None
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            decoded = contents.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValueError("无法识别CSV编码")
    sample = decoded[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    return [dict(row) for row in csv.DictReader(io.StringIO(decoded), dialect=dialect) if any(str(value or "").strip() for value in row.values())]


def _rows_from_xlsx(contents: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - exercised in deployment verification
        raise ValueError("Excel导入组件未安装") from exc
    workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not values:
        return []
    header_index = next((index for index, row in enumerate(values[:20]) if sum(1 for cell in row if str(cell or "").strip() in DATE_HEADERS + AMOUNT_HEADERS + INFLOW_HEADERS + OUTFLOW_HEADERS) >= 2), 0)
    headers = [str(value or "").strip() for value in values[header_index]]
    return [dict(zip(headers, row)) for row in values[header_index + 1:] if any(value not in (None, "") for value in row)]


def _rows_from_pdf(contents: bytes) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(contents)).pages)
    rows: list[dict[str, Any]] = []
    pattern = re.compile(r"(20\d{2}[-/.年]\d{1,2}[-/.月]\d{1,2}日?).*?([+\-]?\d[\d,]*\.\d{2})(?:\s+)([^\n]+)")
    for index, match in enumerate(pattern.finditer(text), 1):
        amount = match.group(2)
        rows.append({
            "交易日期": match.group(1),
            "交易金额": amount,
            "收支类型": "支出" if amount.startswith("-") else "收入",
            "摘要": match.group(3).strip(),
            "交易流水号": f"pdf-row-{index}",
        })
    if not rows:
        raise ValueError("PDF中未识别到结构化流水；扫描版请改用截图识别或导出Excel/CSV")
    return rows


def parse_finance_file(
    filename: str,
    content_type: str,
    contents: bytes,
    *,
    account_owner_kind: str = "unknown",
) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".txt"} or "csv" in content_type:
        rows = _rows_from_csv(contents)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _rows_from_xlsx(contents)
    elif suffix == ".pdf" or content_type == "application/pdf":
        rows = _rows_from_pdf(contents)
    else:
        raise ValueError("当前账单导入支持 CSV、Excel 和文本型 PDF；图片请使用拍照识别")

    transactions: list[dict[str, Any]] = []
    inflow_total = 0
    outflow_total = 0
    balances: list[int] = []
    for index, row in enumerate(rows, 1):
        direction, amount_minor = _direction(row)
        transaction_date = _normalize_date(_first(row, DATE_HEADERS))
        if not transaction_date or amount_minor <= 0 or direction == "unknown":
            continue
        counterparty = _first(row, COUNTERPARTY_HEADERS)
        summary = _first(row, SUMMARY_HEADERS)
        reference = _first(row, REFERENCE_HEADERS) or f"row-{index}"
        classification = classify_finance_transaction(
            direction=direction,
            counterparty=counterparty,
            summary=summary,
            account_owner_kind=account_owner_kind,
        )
        if direction == "inflow":
            inflow_total += amount_minor
        else:
            outflow_total += amount_minor
        balance = _amount(_first(row, BALANCE_HEADERS))
        if balance:
            balances.append(balance)
        transactions.append({
            "transaction_date": transaction_date,
            "transaction_time": _first(row, TIME_HEADERS),
            "direction": direction,
            "amount_minor": amount_minor,
            "counterparty": counterparty,
            "summary": summary,
            "source_reference": reference,
            "raw_data": {str(key): value for key, value in row.items()},
            **classification,
        })
    if not transactions:
        raise ValueError("账单中没有识别到有效的日期、收支方向和金额")
    return {
        "row_count": len(transactions),
        "inflow_minor": inflow_total,
        "outflow_minor": outflow_total,
        "opening_balance_minor": balances[0] if balances else None,
        "closing_balance_minor": balances[-1] if balances else None,
        "transactions": transactions,
    }
